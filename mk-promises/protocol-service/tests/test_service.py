"""Offline tests use real document bytes and mocked Grid HTTP, never Knesset."""

import http.client
import importlib.util
import io
import json
import os
from pathlib import Path
import socket
import sys
import threading
import time
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request

MODULE_PATH = Path(__file__).resolve().parents[1] / "server.py"
spec = importlib.util.spec_from_file_location("protocol_service", MODULE_PATH)
service = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = service
spec.loader.exec_module(service)

PAGE = "https://main.knesset.gov.il/Activity/committees/pages/AllCommitteesAgenda.aspx?Tab=3&ItemID=123"
DOCUMENT = "https://fs.knesset.gov.il/25/Committees/25_ptv_123.doc"
TRANSCRIPT = (
    "פרוטוקול מס׳ 123\nמישיבת ועדת הכלכלה\nשם הוועדה: ועדת הכלכלה\nתאריך: 28.10.2026\n"
    "סדר היום: דיון בתחבורה ציבורית\nנכחו:\nחברי הוועדה:\n"
    + "יושב ראש: נפתח את הדיון. נציגת המשרד: להלן הנתונים שנאספו לגבי השירות לציבור.\n" * 6
)


class Response(io.BytesIO):
    def __init__(self, raw, *, url=DOCUMENT, status=200, headers=None):
        super().__init__(raw)
        self.status, self.url, self.headers = status, url, headers or {}

    def geturl(self):
        return self.url


class FakeGridHTTP:
    def __init__(self, *, snapshot=None, fail=None, final_url=PAGE):
        self.calls, self.fail, self.final_url = [], fail, final_url
        self.snapshot = snapshot or {"finalUrl": PAGE, "title": "פרוטוקול", "text": TRANSCRIPT,
                                     "links": [], "status": 200, "login": False, "captcha": False}

    def __call__(self, request, timeout):
        path = request.full_url.removeprefix("http://grid:4444")
        method = request.get_method()
        payload = json.loads(request.data) if request.data else None
        self.calls.append((method, path, payload))
        if self.fail and self.fail == (method, path):
            raise urllib.error.HTTPError(request.full_url, 500, "PRIVATE SERVER ERROR", {}, io.BytesIO(b"secret"))
        if path == "/session" and method == "POST":
            value = {"sessionId": "test-session", "capabilities": {"browserName": "chrome"}}
        elif path.endswith("/execute/sync"):
            value = self.snapshot
        elif path.endswith("/url") and method == "GET":
            value = self.final_url
        else:
            value = None
        return Response(json.dumps({"value": value}).encode(), url=request.full_url)


def config(**kwargs):
    return service.Config("http://grid:4444", **kwargs)


def docx_fixture():
    from docx import Document
    document = Document()
    document.core_properties.title = "פרוטוקול בדיקה"
    for paragraph in TRANSCRIPT.split("\n"):
        document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "עמודה א"
    table.cell(0, 1).text = "עמודה ב"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_fixture(with_text=True):
    # Minimal valid PDF with xref offsets; fixture construction has no extra dependency.
    content = b"BT /F1 12 Tf 40 700 Td (Known committee fixture text) Tj ET" if with_text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    raw, offsets = b"%PDF-1.4\n", [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw += str(index).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    start = len(raw)
    raw += b"xref\n0 6\n0000000000 65535 f \n"
    for offset in offsets[1:]:
        raw += f"{offset:010} 00000 n \n".encode()
    return raw + b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(start).encode() + b"\n%%EOF\n"


class InputTests(unittest.TestCase):
    def test_official_url_and_default_port(self):
        self.assertEqual(service.official_url("https://FS.KNESSET.GOV.IL:443/x.PDF#page=2"), "https://fs.knesset.gov.il/x.PDF")
        self.assertEqual(service.validate_request(json.dumps({"url": PAGE}).encode()), PAGE)

    def test_external_ambiguous_and_local_urls_rejected(self):
        for url in ["http://fs.knesset.gov.il/x.pdf", "https://knesset.gov.il.evil.test/x", "https://evilknesset.gov.il/x",
                    "https://127.0.0.1/x", "https://knesset.gov.il:8443/x", "https://user@knesset.gov.il/x",
                    "https://knesset.gov.il./x", "https://knesset.gov.il\\@evil.test/x", "https://knesset.gov.il/x\n",
                    "https://-bad.knesset.gov.il/x", "javascript:alert(1)", "file:///etc/passwd"]:
            with self.subTest(url=url), self.assertRaises(service.ServiceError):
                service.official_url(url)

    def test_strict_body_limits_and_no_javascript_input(self):
        for raw in [b'{"url":"a","url":"b"}', b'{"url":NaN}', b'[]', b'{"url":1}',
                    json.dumps({"url": PAGE, "script": "arbitrary code"}).encode()]:
            with self.subTest(raw=raw), self.assertRaises(service.ServiceError):
                service.validate_request(raw)
        with self.assertRaises(service.ServiceError) as error:
            service.validate_request(b" " * (service.MAX_BODY_BYTES + 1))
        self.assertEqual(error.exception.status, 413)

    def test_configuration_requires_grid_and_token_for_public_or_null(self):
        for values in [{}, {"SELENIUM_URL": "http://grid:4444", "PROTOCOL_HOST": "0.0.0.0"},
                       {"SELENIUM_URL": "http://grid:4444", "PROTOCOL_ALLOWED_ORIGINS": "null"}]:
            with patch.dict(os.environ, values, clear=True), self.assertRaises(ValueError):
                service.Config.from_env()
        with patch.dict(os.environ, {"SELENIUM_URL": "http://grid:4444", "PROTOCOL_HOST": "0.0.0.0", "PROTOCOL_API_TOKEN": "secret"}, clear=True):
            self.assertEqual(service.Config.from_env().port, 8788)

    def test_exact_origins_only(self):
        self.assertEqual(service.parse_origins("https://app.test,http://localhost:8000"), frozenset({"https://app.test", "http://localhost:8000"}))
        for origin in ["*", "https://*.test", "https://app.test/path", "https://u:p@app.test"]:
            with self.subTest(origin=origin), self.assertRaises(ValueError):
                service.parse_origins(origin)


class GridTests(unittest.TestCase):
    def read(self, fake):
        with patch.object(service.time, "sleep"):
            return service.read_page(PAGE, config(), grid_factory=lambda url, **kw: service.GridClient(url, open_url=fake, **kw))

    def test_real_w3c_http_contract_fixed_script_and_cleanup(self):
        fake = FakeGridHTTP()
        result = self.read(fake)
        self.assertEqual(result["text"], TRANSCRIPT)
        self.assertEqual(result["kind"], "html")
        self.assertTrue(result["isProtocol"])
        self.assertEqual(result["committee"], "ועדת הכלכלה")
        self.assertEqual(result["date"], "28.10.2026")
        self.assertEqual(fake.calls[0][:2], ("POST", "/session"))
        caps = fake.calls[0][2]["capabilities"]["alwaysMatch"]
        self.assertFalse(caps["acceptInsecureCerts"])
        self.assertIn("--incognito", caps["goog:chromeOptions"]["args"])
        scripts = [payload for method, path, payload in fake.calls if path.endswith("/execute/sync")]
        self.assertTrue(all(x == {"script": service.SNAPSHOT_SCRIPT, "args": []} for x in scripts))
        self.assertEqual(fake.calls[-1][:2], ("DELETE", "/session/test-session"))

    def test_document_links_keep_listing_ineligible(self):
        fake = FakeGridHTTP(snapshot={"finalUrl": PAGE, "title": "דיוני הוועדה", "text": TRANSCRIPT,
            "links": [{"url": DOCUMENT, "title": "הפרוטוקול"}, {"url": DOCUMENT, "title": "כפול"},
                      {"url": "https://evil.test/other.pdf", "title": "חיצוני"}]})
        result = self.read(fake)
        self.assertEqual(result["kind"], "page")
        self.assertFalse(result["isProtocol"])
        self.assertEqual(result["documents"], [{"url": DOCUMENT, "title": "הפרוטוקול", "format": "doc"}])

    def test_ordinary_page_is_never_a_transcript(self):
        fake = FakeGridHTTP(snapshot={"finalUrl": PAGE, "title": "פרוטוקולים", "text": "חיפוש פרוטוקולים ודיוני הוועדה", "links": []})
        result = self.read(fake)
        self.assertFalse(result["isProtocol"])
        self.assertIsNone(result["date"])
        self.assertEqual(result["kind"], "page")

    def test_cleanup_after_navigation_failure_and_private_error_not_exposed(self):
        fake = FakeGridHTTP(fail=("POST", "/session/test-session/url"))
        with self.assertRaises(service.ServiceError) as error:
            self.read(fake)
        self.assertNotIn("PRIVATE", error.exception.message)
        self.assertEqual(fake.calls[-1][:2], ("DELETE", "/session/test-session"))

    def test_external_redirect_stops_before_dom_read_and_closes(self):
        fake = FakeGridHTTP(final_url="https://evil.test/")
        with self.assertRaises(service.ServiceError) as error:
            self.read(fake)
        self.assertEqual(error.exception.code, "redirect_denied")
        self.assertFalse(any(path.endswith("execute/sync") for _method, path, _payload in fake.calls))
        self.assertEqual(fake.calls[-1][0], "DELETE")

    def test_geoblock_captcha_and_login_stop_and_close(self):
        for extra in [{"status": 473}, {"captcha": True}, {"login": True}, {"text": "Access denied"},
                      {"finalUrl": "https://www.knesset.gov.il/maintenance-page-geo"}]:
            snapshot = {"finalUrl": PAGE, "text": TRANSCRIPT, "title": "מקור", "links": [], **extra}
            fake = FakeGridHTTP(snapshot=snapshot)
            with self.subTest(extra=extra), self.assertRaises(service.ServiceError) as error:
                self.read(fake)
            self.assertEqual(error.exception.code, "source_blocked")
            self.assertEqual(fake.calls[-1][0], "DELETE")

    def test_grid_body_limit_and_timeout(self):
        client = service.GridClient("http://grid:4444", open_url=lambda *_a, **_k: Response(b"x" * (service.MAX_DRIVER_BYTES + 1)))
        with self.assertRaises(service.ServiceError):
            client.command("GET", "/status")
        with self.assertRaises(service.ServiceError) as error:
            service.bounded_read(Response(b"abc"), 10, time.monotonic() - 1)
        self.assertEqual(error.exception.status, 504)

    def test_chunked_reader_enforces_limit(self):
        self.assertEqual(len(service.bounded_read(Response(b"a" * 100), 20, time.monotonic() + 10)), 21)


class DocumentTests(unittest.TestCase):
    def test_docx_at_doc_url_detected_and_extracted_in_worker(self):
        raw = docx_fixture()
        self.assertEqual(service.document_kind(DOCUMENT), "doc")
        self.assertEqual(service.detect_document(raw), "docx")
        with patch.object(service, "download_document", return_value=(raw, DOCUMENT)):
            result = service.read_protocol(DOCUMENT, config())
        self.assertEqual(result["kind"], "docx")
        self.assertTrue(result["isProtocol"])
        self.assertIn("יושב ראש: נפתח את הדיון", result["text"])
        self.assertIn("עמודה א\tעמודה ב", result["text"])

    def test_legitimate_pdf_fixture_extraction(self):
        raw = pdf_fixture()
        self.assertEqual(service.detect_document(raw), "pdf")
        result = service.extract_document(raw, "pdf")
        self.assertIn("Known committee fixture text", result["text"])

    def test_empty_pdf_reports_ocr_and_cannot_be_evidence(self):
        extracted = service.extract_document(pdf_fixture(False), "pdf")
        self.assertEqual(extracted["text"].strip(), "")
        self.assertTrue(any("OCR" in item for item in extracted["warnings"]))
        result = service.result_for(DOCUMENT, DOCUMENT, "מסמך", extracted["text"], "pdf", warnings=extracted["warnings"])
        self.assertFalse(result["isProtocol"])

    def test_background_document_never_marked_by_committee_path(self):
        url = "https://fs.knesset.gov.il/25/Committees/25_cs_bg_123.pdf"
        result = service.result_for(url, url, "מסמך רקע", TRANSCRIPT, "pdf")
        self.assertFalse(result["isProtocol"])

    def test_protocol_filename_alone_never_confirms_readable_protocol(self):
        for text in ["corrupt data", "", "unrelated extracted text " * 100, "פרוטוקול מס׳ 1"]:
            result = service.result_for(DOCUMENT, DOCUMENT, "מסמך", text, "pdf")
            self.assertFalse(result["isProtocol"])
        minimal_header = "פרוטוקול מס׳ 1\nמישיבת ועדת הכלכלה\n" + "דברי הדיון כפי שנרשמו. " * 30
        self.assertTrue(service.result_for(DOCUMENT, DOCUMENT, "מסמך", minimal_header, "doc")["isProtocol"])
        other = "https://fs.knesset.gov.il/document.pdf"
        self.assertFalse(service.result_for(other, other, "מסמך", minimal_header, "pdf")["isProtocol"])

    def test_legacy_doc_missing_reader_is_explicit(self):
        with patch.object(service.shutil, "which", return_value=None), self.assertRaises(service.ServiceError) as error:
            service.extract_document(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", "doc")
        self.assertEqual(error.exception.code, "doc_reader_missing")

    def test_unsupported_or_html_error_file_not_passed_off_as_document(self):
        for raw in [b"<html>not a document</html>", b"", b"not-pdf"]:
            with self.subTest(raw=raw), self.assertRaises(service.ServiceError):
                service.detect_document(raw)

    def test_unicode_length_and_bounded_response(self):
        text = "😀" * (service.MAX_TEXT_CHARS + 1)
        result = service.result_for(DOCUMENT, DOCUMENT, "title", text, "pdf", documents=[
            {"url": "https://fs.knesset.gov.il/x.pdf?q=" + "x" * 4000, "title": "a" * 300, "format": "pdf"}
            for _ in range(50)])
        self.assertTrue(result["truncated"])
        self.assertEqual(len(result["text"]), service.MAX_TEXT_CHARS)
        self.assertLessEqual(len(service.json_bytes(result)), service.MAX_RESPONSE_BYTES)

    def test_download_external_redirect_is_not_followed(self):
        calls = []
        def opener(request, **_kwargs):
            calls.append(request.full_url)
            raise urllib.error.HTTPError(request.full_url, 302, "redirect", {"Location": "https://evil.test/x"}, io.BytesIO())
        with self.assertRaises(service.ServiceError) as error:
            service.download_document(DOCUMENT, open_url=opener)
        self.assertEqual(error.exception.code, "redirect_denied")
        self.assertEqual(calls, [DOCUMENT])

    def test_download_allowed_redirect_rechecked(self):
        target = "https://fs.knesset.gov.il/25/Committees/other.docx"
        calls = []
        def opener(request, **_kwargs):
            calls.append(request.full_url)
            if len(calls) == 1:
                raise urllib.error.HTTPError(request.full_url, 302, "redirect", {"Location": target}, io.BytesIO())
            return Response(b"fixture", url=target)
        self.assertEqual(service.download_document(DOCUMENT, open_url=opener), (b"fixture", target))

    def test_download_limits_and_access_denial(self):
        cases = [Response(b"", headers={"Content-Length": str(service.MAX_DOWNLOAD_BYTES + 1)}),
                 Response(b"access denied", status=473), Response(b"<html>access denied</html>")]
        for response in cases:
            with self.subTest(status=response.status), self.assertRaises(service.ServiceError) as error:
                service.download_document(DOCUMENT, open_url=lambda *_a, **_kw: response)
            self.assertIn(error.exception.status, {403, 413})


class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = config(api_token="test-token", allowed_origins=frozenset({"https://app.test"}))
        try:
            cls.server = service.ProtocolHTTPServer(("127.0.0.1", 0), cls.config,
                reader=lambda url, _config: service.result_for(url, url, "fixture", TRANSCRIPT, "html"))
        except PermissionError as error:
            raise unittest.SkipTest("Loopback socket creation is unavailable in this sandbox") from error
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=5)

    def request(self, method="POST", body=None, headers=None, path="/api/protocols"):
        body = json.dumps({"url": PAGE}).encode() if body is None else body
        conn = http.client.HTTPConnection(*self.server.server_address, timeout=5)
        try:
            conn.request(method, path, body=body, headers={"Content-Type": "application/json", "Authorization": "Bearer test-token", **(headers or {})})
            response = conn.getresponse()
            raw = response.read()
            return response.status, dict(response.getheaders()), json.loads(raw) if raw else None
        finally:
            conn.close()

    def test_native_without_origin_and_browser_exact_origin(self):
        for headers in [{}, {"Origin": "https://app.test"}]:
            status, response_headers, value = self.request(headers=headers)
            self.assertEqual(status, 200)
            self.assertEqual(value["text"], TRANSCRIPT)
            self.assertEqual(response_headers["Cache-Control"], "no-store")
            if headers:
                self.assertEqual(response_headers["Access-Control-Allow-Origin"], "https://app.test")

    def test_auth_origin_and_null_denied(self):
        for headers, expected in [({"Authorization": "Bearer wrong"}, 401), ({"Origin": "https://evil.test"}, 403), ({"Origin": "null"}, 403)]:
            self.assertEqual(self.request(headers=headers)[0], expected)

    def test_preflight(self):
        status, headers, _ = self.request("OPTIONS", body=b"", headers={"Origin": "https://app.test", "Access-Control-Request-Method": "POST", "Access-Control-Request-Headers": "authorization,content-type", "Authorization": ""})
        self.assertEqual(status, 204)
        self.assertEqual(headers["Access-Control-Allow-Origin"], "https://app.test")

    def test_bad_method_route_media_body_and_url(self):
        self.assertEqual(self.request("GET")[0], 405)
        self.assertEqual(self.request(path="/wrong")[0], 404)
        self.assertEqual(self.request(headers={"Content-Type": "text/plain"})[0], 415)
        self.assertEqual(self.request(body=b"x" * (service.MAX_BODY_BYTES + 1))[0], 413)
        self.assertEqual(self.request(body=b'{"url":"https://evil.test"}')[0], 400)

    def test_two_reader_slots(self):
        self.server.read_slots.acquire()
        self.server.read_slots.acquire()
        try:
            self.assertEqual(self.request()[0], 503)
        finally:
            self.server.read_slots.release()
            self.server.read_slots.release()


@unittest.skipUnless(os.environ.get("SELENIUM_TEST_URL"), "Set SELENIUM_TEST_URL for an actual Grid smoke test")
class RealGridSmokeTest(unittest.TestCase):
    def test_fresh_browser_fixed_dom_and_delete_session(self):
        calls = []
        opener = urllib.request.build_opener(service.NoRedirect()).open
        def tracked(request, **kwargs):
            calls.append((request.get_method(), request.full_url))
            return opener(request, **kwargs)
        grid = service.GridClient(os.environ["SELENIUM_TEST_URL"], open_url=tracked)
        sid = None
        try:
            sid = grid.new_session()
            # Test-only low-level use; production read_protocol still only accepts Knesset HTTPS.
            grid.session_command("POST", "/url", {"url": "about:blank"})
            text = grid.session_command("POST", "/execute/sync", {
                "script": "document.title='Protocol fixture'; document.body.innerText=arguments[0]; return document.body.innerText;",
                "args": [TRANSCRIPT],
            })
            self.assertEqual(text, TRANSCRIPT)
            self.assertEqual(grid.session_command("GET", "/url"), "about:blank")
        finally:
            grid.close()
        self.assertIsNone(grid.session_id)
        self.assertEqual(calls[-1], ("DELETE", os.environ["SELENIUM_TEST_URL"].rstrip("/") + "/session/" + sid))
        with self.assertRaises(service.ServiceError):
            grid.command("GET", "/session/" + sid + "/url")


if __name__ == "__main__":
    unittest.main()
